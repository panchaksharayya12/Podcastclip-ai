import sys
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.text import PP_ALIGN

def set_cell_background(cell, fill_hex):
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tc_pr.append(shd)

def create_word_report(output_path):
    print("Generating Word Report without emojis...")
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Document Title
    title = doc.add_paragraph()
    title_run = title.add_run("CLIPFORGE AI")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(28)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(17, 19, 24)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(2)

    # Subtitle
    sub = doc.add_paragraph()
    sub_run = sub.add_run('Comprehensive Full-Stack Project Report\n"One conversation. A hundred moments worth sharing."')
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(90, 95, 105)
    sub.paragraph_format.space_after = Pt(16)

    # Metadata Block
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(20)
    meta_runs = [
        ("Live Deployment: ", True),
        ("https://podcastclip-ai.vercel.app\n", False),
        ("GitHub Repository: ", True),
        ("https://github.com/panchaksharayya12/Podcastclip-ai\n", False),
        ("Architecture: ", True),
        ("React, TypeScript, Vite, Tailwind CSS, Express, Node.js, FFmpeg\n", False),
        ("Date: ", True),
        ("October 2026\n", False),
        ("Author: ", True),
        ("ClipForge AI Engineering Team\n", False),
    ]
    for text, bold in meta_runs:
        r = meta_p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.bold = bold
        if "https://" in text:
            r.font.color.rgb = RGBColor(0, 102, 204)
        else:
            r.font.color.rgb = RGBColor(50, 50, 50)

    doc.add_heading("1. Executive Summary", level=1)
    p = doc.add_paragraph(
        "ClipForge AI is an end-to-end, full-stack video intelligence platform that bridges long-form content creation "
        "and short-form social distribution. Modern content creators, podcast hosts, interviewers, and educators "
        "frequently produce rich 45-to-90 minute conversations but face prohibitive friction when extracting concise, "
        "engaging segments optimized for algorithmic mobile feeds such as YouTube Shorts, Instagram Reels, TikTok, and LinkedIn."
    )
    p.paragraph_format.space_after = Pt(8)

    p2 = doc.add_paragraph(
        "The ClipForge AI system addresses this pipeline end-to-end: from multi-format media ingestion (MP4, MOV, WebM, MP3, WAV) "
        "and automated duration/stream probing, through deterministic transcript heuristic parsing and OpenAI GPT-4o-mini "
        "structured moment detection, to authentic server-side FFmpeg rendering of 9:16 vertical MP4 video clips with "
        "cinematic blurred background padding."
    )
    p2.paragraph_format.space_after = Pt(14)

    doc.add_heading("2. Product Vision and Core Requirements", level=1)
    p = doc.add_paragraph(
        "The project was engineered around strict production guidelines:\n"
        "- Authentic Video Processing: Video clips are rendered as genuine H.264/AAC MP4 files via FFmpeg, not mock placeholders.\n"
        "- Dual Intelligence Engines: Full functionality in offline mode via a deterministic sentence/hook parser, with optional OpenAI Whisper and GPT-4o-mini structured analysis.\n"
        "- Transparent Potential Scoring: Engagement indicators are clearly labeled as qualitative AI estimates rather than simulated real-time platform analytics.\n"
        "- Data Privacy: Media processing remains on the local development server without unauthorized third-party telemetry.\n"
        "- Full Responsiveness: Complete mobile, tablet, and desktop compatibility with smooth motion graphics."
    )
    p.paragraph_format.space_after = Pt(14)

    doc.add_heading("3. System Architecture", level=1)
    p = doc.add_paragraph(
        "The platform follows a decoupled client-server architecture with clear segregation between presentation, "
        "route controllers, processing pipelines, and local disk persistence."
    )
    p.paragraph_format.space_after = Pt(8)

    # Architecture Table
    table = doc.add_table(rows=6, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Layer", "Technology", "Primary Responsibilities"]
    col_widths = [Inches(1.5), Inches(1.8), Inches(3.2)]

    for idx, heading in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = heading
        set_cell_background(cell, "191B21")
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)

    rows_data = [
        ("Frontend UI", "React 18, TypeScript, Tailwind CSS", "Editorial interface, motion reveals, video playback controls, caption customization"),
        ("Motion Engine", "Motion for React, HTML5 Canvas", "Word-by-word headline reveals, audio waveform visualizations, staged loading"),
        ("API Gateway", "Express, Node.js, TypeScript", "REST endpoints, CORS security, HTTP 206 video range streaming, input validation"),
        ("Video Pipeline", "FFmpeg (ffmpeg-static, fluent-ffmpeg)", "Probing codecs, audio extraction, 9:16 vertical blurred crop, H.264 encoding"),
        ("AI / Heuristics", "OpenAI API & Deterministic Engine", "Whisper transcription, hook detection, rhetorical question scoring, SRT timing"),
    ]

    for row_idx, data in enumerate(rows_data, start=1):
        for col_idx, text in enumerate(data):
            cell = table.cell(row_idx, col_idx)
            cell.text = text
            set_cell_background(cell, "F9F9F9" if row_idx % 2 == 1 else "FFFFFF")
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(30, 30, 30)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    doc.add_heading("4. Functional Modules", level=1)

    modules = [
        ("4.1 Cinematic Hero and Motion Graphics",
         "Features an abstract HTML5 canvas waveform visualization rendering sine curves and audio spectrum bars. "
         "The headline 'Your best moments deserve the spotlight' uses staggered word reveals respecting user accessibility settings."),
        ("4.2 Dual-Input Studio Workspace",
         "Allows creators to either upload audio/video files with real-time drag-and-drop progress, or paste/import text and SRT transcripts. "
         "A preloaded demo dataset provides immediate access to test the pipeline without uploading external assets."),
        ("4.3 Discovery Parameters and Heuristic Engine",
         "Users configure target durations (15-30s, 30-60s, 60-90s), target platforms, and content archetypes. "
         "The deterministic engine isolates sentence boundaries, opening challenge hooks, and question-answer rhythms."),
        ("4.4 Real Video Rendering Pipeline",
         "The FFmpeg service scales widescreen videos into 9:16 vertical outputs using a dual-stream filter graph: "
         "a background stream zoomed and blurred with boxblur to fill the frame, overlaid by the sharp centered foreground video."),
        ("4.5 Animated Caption Typography Studio",
         "Offers a live vertical smartphone preview with adjustable typography sizes, electric lime keyword highlighting, "
         "safe margin calculations, and one-click standard SRT file generation."),
        ("4.6 Project Library and Universal Exporters",
         "Maintains project metadata and suggested clips in persistent disk storage with localStorage fallback. "
         "Users can export clips as CSV spreadsheets, JSON payloads, plain text excerpts, or standalone MP4 video files.")
    ]

    for mod_title, mod_desc in modules:
        doc.add_heading(mod_title, level=2)
        p = doc.add_paragraph(mod_desc)
        p.paragraph_format.space_after = Pt(8)

    doc.add_heading("5. Verification and Automated Test Results", level=1)
    p = doc.add_paragraph(
        "A rigorous 44-assertion automated test suite was constructed and executed against the live application. "
        "All test categories passed with zero errors:"
    )
    p.paragraph_format.space_after = Pt(6)

    # Test Results Table
    test_table = doc.add_table(rows=7, cols=3)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    test_headers = ["Test Category", "Assertions", "Status"]

    for idx, heading in enumerate(test_headers):
        cell = test_table.cell(0, idx)
        cell.text = heading
        set_cell_background(cell, "191B21")
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)

    test_rows = [
        ("Server Health & Engine Diagnostics", "GET /api/health returns 200, FFmpeg verified, OpenAI flag verified", "PASSED"),
        ("Deterministic Transcript Analysis", "Boundary detection, non-negative start/end, hooks, captions, tags", "PASSED"),
        ("Validation & Error Resilience", "Rejection of empty, malformed, or truncated inputs with HTTP 400", "PASSED"),
        ("Captions & Multi-Format Exporters", "Valid SRT timecodes, CSV spreadsheet headers, JSON schemas", "PASSED"),
        ("Project Persistence CRUD", "Save, list, fetch by unique ID, and clean deletion verified", "PASSED"),
        ("FFmpeg Video Processing Pipeline", "Synthetic video generation, upload probing, 9:16 render, MP4 stream", "PASSED"),
    ]

    for row_idx, data in enumerate(test_rows, start=1):
        for col_idx, text in enumerate(data):
            cell = test_table.cell(row_idx, col_idx)
            cell.text = text
            set_cell_background(cell, "F9F9F9" if row_idx % 2 == 1 else "FFFFFF")
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    if col_idx == 2:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0, 130, 0)
                    else:
                        run.font.color.rgb = RGBColor(30, 30, 30)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    doc.add_heading("6. Deployment and Production Links", level=1)
    p = doc.add_paragraph(
        "The application has been deployed to production on Vercel and pushed to GitHub with clean documentation:\n"
        "- Production Deployment: https://podcastclip-ai.vercel.app\n"
        "- GitHub Repository: https://github.com/panchaksharayya12/Podcastclip-ai\n"
        "- Local Development Port: http://localhost:5173\n"
        "- Backend API Port: http://localhost:5000"
    )
    p.paragraph_format.space_after = Pt(14)

    doc.save(output_path)
    print(f"Word Report created successfully at: {output_path}")

def create_powerpoint_presentation(output_path):
    print("Generating PowerPoint Presentation without emojis...")
    prs = Presentation()
    prs.slide_width = PptxInches(13.333)
    prs.slide_height = PptxInches(7.5)

    blank_layout = prs.slide_layouts[6]

    BG_COLOR = PptxRGBColor(8, 9, 11)        # #08090B
    PANEL_COLOR = PptxRGBColor(25, 27, 33)   # #191B21
    TEXT_WHITE = PptxRGBColor(255, 255, 255)
    TEXT_MUTED = PptxRGBColor(167, 171, 180) # #A7ABB4
    ACCENT_LIME = PptxRGBColor(199, 243, 107)# #C7F36B

    def add_slide_background(slide):
        shape = slide.shapes.add_shape(
            1, 0, 0, prs.slide_width, prs.slide_height # 1 is MSO_SHAPE.RECTANGLE
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = BG_COLOR
        shape.line.fill.background()

    slides_content = [
        # Slide 1: Title
        {
            "badge": "AI-POWERED SHORTS STUDIO",
            "title": "CLIPFORGE AI",
            "subtitle": "One conversation. A hundred moments worth sharing.",
            "bullets": [
                "Production URL: https://podcastclip-ai.vercel.app",
                "Repository: https://github.com/panchaksharayya12/Podcastclip-ai",
                "Tech Stack: React, TypeScript, Vite, Tailwind CSS, Express, Node.js, FFmpeg",
                "Full-Stack Architecture & Project Presentation"
            ]
        },
        # Slide 2: The Problem
        {
            "badge": "INDUSTRY CHALLENGE",
            "title": "The Long-Form Creator Dilemma",
            "subtitle": "Valuable dialogue is trapped inside lengthy recordings",
            "bullets": [
                "Podcasters and interviewers produce 60+ minute episodes with high depth.",
                "Mobile algorithms reward 15-to-60 second vertical clips with immediate hooks.",
                "Manual scrubbing, timestamp logging, and subtitle synchronization take hours.",
                "Existing SaaS tools charge high subscriptions and push generic watermarks."
            ]
        },
        # Slide 3: The Solution
        {
            "badge": "PRODUCT VISION",
            "title": "ClipForge AI: Automated Moment Discovery",
            "subtitle": "From long recordings to distribution-ready short video clips",
            "bullets": [
                "Automated ingestion of MP4, MOV, WebM, MP3, WAV, and M4A media.",
                "Intelligent moment boundary detection tailored for Shorts, Reels, and TikTok.",
                "Dual analysis modes: OpenAI GPT-4o-mini and an offline deterministic heuristic engine.",
                "Authentic server-side video rendering with FFmpeg (9:16 vertical, 1:1, 16:9)."
            ]
        },
        # Slide 4: System Architecture
        {
            "badge": "ENGINEERING DESIGN",
            "title": "Full-Stack System Architecture",
            "subtitle": "Decoupled client-server design for performance and privacy",
            "bullets": [
                "Frontend: React 18 with TypeScript, Vite bundler, and Tailwind CSS palette.",
                "Motion Layer: Motion for React with fluid typography reveals and canvas waveforms.",
                "Backend: Node.js Express server with HTTP 206 range streaming for video seeking.",
                "Local Storage: Secure server-side uploads and generated clips directory."
            ]
        },
        # Slide 5: Core Features
        {
            "badge": "STUDIO CAPABILITIES",
            "title": "Comprehensive Feature Matrix",
            "subtitle": "Engineered for speed, control, and multi-platform publishing",
            "bullets": [
                "Dual Input Cards: Media upload with progress probe + Transcript text/SRT parser.",
                "Discovery Parameters: Duration targets (15-30s, 30-60s, 60-90s) and audience styles.",
                "Interactive Video Preview: HTML5 player seeking exact timestamps with loop toggle.",
                "Universal Exports: Direct MP4 downloads, SRT subtitle tracks, and CSV/JSON metadata."
            ]
        },
        # Slide 6: Real FFmpeg Pipeline
        {
            "badge": "VIDEO PROCESSING",
            "title": "Authentic FFmpeg Rendering Pipeline",
            "subtitle": "Real H.264 / AAC video clipping with cinematic framing",
            "bullets": [
                "Bundled ffmpeg-static provides a standalone executable without external dependencies.",
                "Vertical 9:16 Conversion: Dual-stream filter graph zooms & blurs background fill.",
                "Zero Image Distortion: Sharp centered video preserves original aspect ratio.",
                "Optimized Web Playback: Encoded with faststart flags for rapid browser streaming."
            ]
        },
        # Slide 7: Animated Captions & Typography
        {
            "badge": "MOBILE PRESENTATION",
            "title": "Animated Caption Studio",
            "subtitle": "Social-first kinetic subtitle styling and timing",
            "bullets": [
                "Live 9:16 vertical mobile device simulator displaying real-time styling.",
                "Typography controls: Font sizing (small, medium, large) and electric lime highlighting.",
                "Background pill and outline styling with vertical safe margins for mobile UI.",
                "Standard SRT export ready to import into Premiere Pro, Final Cut, and CapCut."
            ]
        },
        # Slide 8: Verification & Testing
        {
            "badge": "QUALITY ASSURANCE",
            "title": "Rigorous Verification: 44 Test Assertions",
            "subtitle": "Complete automated testing covering API, heuristics, and video generation",
            "bullets": [
                "Health Diagnostics: Verified live server status, FFmpeg detection, and memory limits.",
                "Deterministic Engine: Tested sentence boundaries, hooks, and score breakdowns.",
                "Validation Testing: Rejected empty and malformed transcripts with HTTP 400.",
                "Video Rendering: Generated synthetic MP4, uploaded, trimmed, and downloaded real clip."
            ]
        },
        # Slide 9: Deployment & Access
        {
            "badge": "PRODUCTION LAUNCH",
            "title": "Deployment & Live Environments",
            "subtitle": "Available online and structured for simple local extension",
            "bullets": [
                "Production Website: https://podcastclip-ai.vercel.app",
                "GitHub Source Code: https://github.com/panchaksharayya12/Podcastclip-ai",
                "Single Command Startup: 'npm run dev' boots backend and frontend concurrently.",
                "Zero Required API Keys: Runs in complete offline mode out of the box."
            ]
        }
    ]

    for slide_data in slides_content:
        slide = prs.slides.add_slide(blank_layout)
        add_slide_background(slide)

        # Panel Card
        panel = slide.shapes.add_shape(
            1, PptxInches(0.8), PptxInches(0.8), PptxInches(11.733), PptxInches(5.9)
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = PANEL_COLOR
        panel.line.fill.background()

        # Text Frame
        tb = slide.shapes.add_textbox(
            PptxInches(1.2), PptxInches(1.2), PptxInches(10.9), PptxInches(5.1)
        )
        tf = tb.text_frame
        tf.word_wrap = True

        # Badge
        p_badge = tf.paragraphs[0]
        p_badge.text = slide_data["badge"]
        p_badge.font.size = PptxPt(11)
        p_badge.font.bold = True
        p_badge.font.color.rgb = ACCENT_LIME
        p_badge.space_after = PptxPt(8)

        # Title
        p_title = tf.add_paragraph()
        p_title.text = slide_data["title"]
        p_title.font.size = PptxPt(32)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.space_after = PptxPt(6)

        # Subtitle
        p_sub = tf.add_paragraph()
        p_sub.text = slide_data["subtitle"]
        p_sub.font.size = PptxPt(16)
        p_sub.font.italic = True
        p_sub.font.color.rgb = TEXT_MUTED
        p_sub.space_after = PptxPt(20)

        # Bullets
        for b_text in slide_data["bullets"]:
            p_bullet = tf.add_paragraph()
            p_bullet.text = "•  " + b_text
            p_bullet.font.size = PptxPt(14)
            p_bullet.font.color.rgb = PptxRGBColor(220, 225, 230)
            p_bullet.space_after = PptxPt(10)

    prs.save(output_path)
    print(f"PowerPoint Presentation created successfully at: {output_path}")

if __name__ == "__main__":
    report_file = os.path.abspath("CLIPFORGE_AI_PROJECT_REPORT.docx")
    ppt_file = os.path.abspath("CLIPFORGE_AI_PRESENTATION.pptx")
    create_word_report(report_file)
    create_powerpoint_presentation(ppt_file)
